"""生成本次代码修正说明、操作手册和脚本流程图。"""

from pathlib import Path

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from PIL import Image, ImageDraw, ImageFont
except ModuleNotFoundError as exc:
    raise SystemExit(
        "缺少生成交付文档所需依赖，请先运行："
        "pip install -r tools/requirements-docs.txt") from exc


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "非代码资料" / "交付文档"
FLOW_PATH = OUT_DIR / "重启测试脚本逻辑流程图.png"
CHANGE_DOC = OUT_DIR / "设备重启测试脚本本次修改说明.docx"
GUIDE_DOC = OUT_DIR / "设备重启测试傻瓜操作手册.docx"

FONT_PATHS = [
    Path(r"C:\Windows\Fonts\msyh.ttc"),
    Path(r"C:\Windows\Fonts\simhei.ttf"),
]


def _font(size, bold=False):
    for path in FONT_PATHS:
        if path.exists():
            return ImageFont.truetype(str(path), size, index=0)
    return ImageFont.load_default()


def _set_run_font(run, size=10.5, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def _shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _style_doc(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.3

    for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(14 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)
        if style_name == "Title":
            # Word 内置 Title 样式可能带蓝色下划线，交付文档不使用标题分隔线。
            p_pr = style.element.get_or_add_pPr()
            for border in p_pr.findall(qn("w:pBdr")):
                p_pr.remove(border)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("设备重启测试脚本")
    _set_run_font(run, 8, color=(100, 100, 100))


def _title(doc, title, subtitle):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(p.add_run(title), 22, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(p.add_run(subtitle), 10, color=(90, 90, 90))


def _paragraph(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        _set_run_font(p.add_run(bold_lead), 10.5, bold=True)
        _set_run_font(p.add_run(text[len(bold_lead):]), 10.5)
    else:
        _set_run_font(p.add_run(text), 10.5)
    return p


def _bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    _set_run_font(p.add_run(text), 10.5)
    return p


def _number(doc, text):
    p = doc.add_paragraph(style="List Number")
    _set_run_font(p.add_run(text), 10.5)
    return p


def _heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    _set_run_font(p.add_run(text), 15 if level == 1 else 12, bold=True)
    return p


def _table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        _shade(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_font(p.add_run(header), 9.5, bold=True, color=(255, 255, 255))
        if widths:
            cell.width = Cm(widths[i])
    # 表格跨页时重复表头，避免读者在第二页看不懂每列含义。
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for r, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            if r % 2 == 1:
                _shade(cell, "F3F7FA")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _set_run_font(p.add_run(str(value)), 9.2)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph()
    return table


def _draw_centered(draw, box, text, font, fill, outline=(65, 95, 120)):
    x1, y1, x2, y2 = box
    draw.rounded_rectangle(box, radius=18, fill=f"#{fill}", outline=outline, width=3)
    left, top, right, bottom = draw.multiline_textbbox((0, 0), text, font=font, spacing=7, align="center")
    draw.multiline_text(((x1 + x2 - (right - left)) / 2, (y1 + y2 - (bottom - top)) / 2),
                        text, font=font, fill=(25, 35, 45), spacing=7, align="center")


def _arrow(draw, x1, y1, x2, y2):
    draw.line((x1, y1, x2, y2), fill=(45, 95, 135), width=5)
    if y2 > y1:
        draw.polygon([(x2, y2), (x2 - 10, y2 - 17), (x2 + 10, y2 - 17)], fill=(45, 95, 135))
    else:
        draw.polygon([(x2, y2), (x2 - 10, y2 + 17), (x2 + 10, y2 + 17)], fill=(45, 95, 135))


def create_flowchart():
    image = Image.new("RGB", (1500, 1710), "white")
    draw = ImageDraw.Draw(image)
    title_font = _font(42, bold=True)
    box_font = _font(27)
    small_font = _font(23)
    draw.text((430, 42), "设备重启测试脚本逻辑流程", font=title_font, fill=(25, 50, 80))

    boxes = [
        (430, 130, 1070, 245, "运行 python main.py\n先检查设备配置是否正确"),
        (430, 310, 1070, 425, "配置无误后确认开始\n为每台设备创建同批次报告文件夹"),
        (430, 490, 1070, 605, "每台设备启动一个测试线程\n打开串口并写入日志"),
        (430, 670, 1070, 795, "发送准备命令和 reboot\n开始等待启动关键词"),
        (430, 860, 1070, 985, "同时找到两个启动关键词？\niot online 和 ==>trim:"),
        (430, 1050, 1070, 1175, "成功：记 PASS\n超时：补发 reboot，达到上限记 FAIL\n串口异常：记 ERROR"),
        (430, 1240, 1070, 1355, "完成计划轮次或收到 Ctrl+C\n保存 Excel、摘要、完整日志"),
        (430, 1420, 1070, 1535, "主线程只汇总本次文件\n生成文本汇总和 HTML 报告"),
    ]
    colors = ["DCEAF7", "E6F1FA", "EAF4EA", "FFF2CC", "FCE4D6", "F4CCCC", "E2F0D9", "D9EAD3"]
    for box, color in zip(boxes, colors):
        _draw_centered(draw, box[:4], box[-1], box_font, color)
    for i in range(len(boxes) - 1):
        _arrow(draw, 750, boxes[i][3] + 10, 750, boxes[i + 1][1] - 10)

    # 为判断节点补充说明，避免只看箭头时误解。
    draw.rounded_rectangle((1095, 855, 1440, 990), radius=14, fill="#FFF2CC", outline=(160, 120, 25), width=2)
    draw.multiline_text((1120, 875), "按 Ctrl+C：\n取消命令等待和重连\n先保存已完成结果\n再关闭串口", font=small_font, fill=(75, 60, 20), spacing=7)
    draw.line((1070, 925, 1090, 925), fill=(160, 120, 25), width=3)
    image.save(FLOW_PATH)


def create_change_doc():
    doc = Document()
    _style_doc(doc)
    _title(doc, "设备重启测试脚本本次修改说明", "说明已修正的位置、原因、实际作用和验证结果")
    _paragraph(doc, "这次修改的目标是让测试结果更可信、报告只使用本次数据，并在用户停止测试时安全地收尾。没有改动设备清单内容、重启次数、启动关键词或超时时间。")

    _heading(doc, "一 修改结果概览")
    _bullet(doc, "成功率只看真正测到的 PASS 和 FAIL，不再把串口或脚本异常当成设备失败。")
    _bullet(doc, "设备串口打不开时，文本汇总和 HTML 报告都会明确显示 ERROR，不再被静默忽略。")
    _bullet(doc, "启动前会检查设备名称、COM 口和 SN 是否重复，避免两台设备抢同一份报告或同一个串口。")
    _bullet(doc, "HTML 报告固定读取本次运行刚生成的文件，不会在本次文件缺失时偷偷拿历史测试结果来凑。")
    _bullet(doc, "报告保存期间产生的最后几条日志会写进日志文件，方便事后排查。")
    _bullet(doc, "按 Ctrl+C 后会通知正在运行的设备线程停止并正常关闭串口，而不是依赖后台线程被强行结束。")

    _heading(doc, "二 每一处改动和作用")
    rows = [
        ("utils/helpers.py", "calc_success_rate", "以前 ERROR 也会拉低成功率。", "成功率改为 PASS ÷ (PASS + FAIL)。ERROR 单独展示，表示本次没有测成。"),
        ("reports/report_generator.py", "设备级错误汇总", "串口在第一轮前就打不开时，原来会被当作执行了一次重启。", "改为设备级错误：会被报告出来，但不冒充一次实际重启。"),
        ("utils/file_manager.py", "配置检查和文件夹时间", "设备名称、COM 口或 SN 重复时，后创建的目录会覆盖前一个记录。", "开始测试前报出重复项；同一批次的设备文件夹使用同一个时间戳。"),
        ("main.py", "配置拦截和报告输入", "配置问题要等到线程启动后才暴露；HTML 可能读取到历史批次。", "配置不合法立即停止；HTML 明确传入本次 Excel 文件路径。"),
        ("reports/html_report.py", "读取本次文件和设备摘要", "空 Excel 或本次文件缺失时，报告可能忽略设备，或改用旧结果。", "本次运行严格只读传入文件；有摘要错误时仍显示 1 个 ERROR。"),
        ("core/test_runner.py", "保存顺序和停止信号", "原来先关闭日志文件，再保存报告，最后阶段日志可能丢失。", "先保存 Excel、摘要和结束日志，最后关闭日志与串口；并在轮次间检查停止请求。"),
        ("core/serial_manager.py", "命令等待和串口重连", "等待启动之外，命令等待和每次重连还会继续睡眠，Ctrl+C 不能及时退出。", "新增 stop_event；中断后立即取消命令等待和重连重试。"),
        ("core/test_runner.py", "停止检查", "停止请求可能发生在准备命令之后，仍会继续发 reboot。", "每个准备步骤后检查停止请求，收到后不再发送 reboot。"),
        ("tools/requirements-docs.txt", "交付文档依赖", "重新生成 Word 文档时缺少 python-docx 或 Pillow 会直接报导入错误。", "单独列出文档依赖，并在生成脚本中提示安装命令。"),
    ]
    _table(doc, ["文件", "改动位置", "原来的问题", "现在的作用"], rows, [3.0, 3.0, 5.0, 6.3])

    _heading(doc, "三 结果口径 用简单话说明")
    _paragraph(doc, "PASS 的意思是设备在规定时间内启动成功。FAIL 的意思是串口正常、脚本也正常，但设备没有按要求启动。ERROR 的意思是这次根本没有测成，例如串口打不开、串口断开或用户中断。")
    _paragraph(doc, "因此成功率只比较 PASS 和 FAIL。举例：10 次成功、1 次启动失败、2 次串口异常，成功率是 10 ÷ (10 + 1) = 90.91%。那 2 次 ERROR 会在报告里单独显示，提醒你先检查串口或电脑环境。")

    _heading(doc, "四 中断时脚本现在怎么做")
    _number(doc, "用户按 Ctrl+C。")
    _number(doc, "主程序向每台设备线程发送停止通知。")
    _number(doc, "正在等待启动、等待准备命令或等待串口重连的设备会立即结束等待；准备阶段收到停止请求后不再发送 reboot。")
    _number(doc, "每台设备保存已经完成的 Excel、摘要和日志，然后关闭串口。")
    _number(doc, "本次不会生成总汇总报告，避免把未完整执行的结果当成完整测试。")

    _heading(doc, "五 脚本逻辑流程图")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FLOW_PATH), width=Cm(12.7))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(p.add_run("图 1 设备重启测试的主要执行流程"), 9, color=(80, 80, 80))

    _heading(doc, "六 已完成的验证")
    _bullet(doc, "已运行 Python 语法编译检查，当前代码可以导入和编译。")
    _bullet(doc, "已验证成功率示例：1 次 PASS 和 1 次 ERROR 的成功率为 100%，不会误显示 50%。")
    _bullet(doc, "已验证：设备级错误写进摘要后，HTML 数据读取会显示 ERROR。")
    _bullet(doc, "已验证：模拟串口打开失败时，Excel、摘要和结束日志会先保存，再关闭文件。")
    _bullet(doc, "已验证：重复 name、port、sn 会在启动前被拦截；同一批次目录使用同一个时间戳。")
    _bullet(doc, "已验证：传入本次 Excel 文件时，HTML 不会扫描历史目录；缺少 Excel 但有摘要错误时会保留 ERROR。")
    _bullet(doc, "已验证：停止信号已设置时，串口重连会立即返回，不会等待重试间隔。")
    _paragraph(doc, "说明：上述为不连接设备的程序验证。本次没有向真实设备发送 reboot 命令。", bold_lead="说明：")
    doc.save(CHANGE_DOC)


def create_guide_doc():
    doc = Document()
    _style_doc(doc)
    _title(doc, "设备重启测试傻瓜操作手册", "从连接设备到查看报告的逐步说明")
    _paragraph(doc, "照着本文一步一步做即可。开始前先确认串口和设备连接正确；不要在设备正在重启时拔电或关掉电脑。")

    _heading(doc, "一 这个脚本是做什么的")
    _paragraph(doc, "它会通过电脑的串口让设备反复重启，等待设备打印出两个启动成功标志，再把每次结果保存为 Excel、文本和 HTML 报告。它不会自动替你接线，也不能修复设备本身的硬件故障。")

    _heading(doc, "二 开始前要准备什么")
    _table(doc, ["要准备的东西", "你要确认什么"], [
        ("Windows 电脑", "已经安装 Python，并能打开 PowerShell 或命令提示符。"),
        ("USB 转串口线", "设备已经通电，串口线插好，电脑设备管理器里能看到 COM 端口。"),
        ("设备信息", "知道每台设备的 COM 口、波特率、SN、user_id、license_id。"),
        ("项目文件", "文件夹中有 main.py、config、core、reports、utils。"),
    ], [5.0, 12.3])
    _paragraph(doc, "重要提醒：同一个 COM 口不能同时被两个软件占用。请先关闭串口助手、IDE 串口插件和其他测试程序。", bold_lead="重要提醒：")

    _heading(doc, "三 第一次使用 只需要做一次")
    _number(doc, "打开 PowerShell。")
    _number(doc, "进入脚本所在目录，例如：cd E:\\test_scripts\\device_reboot_test")
    _number(doc, "安装依赖：pip install -r requirements.txt")
    _number(doc, "执行检查：python -m compileall -q .。没有报错即可。")

    _heading(doc, "四 填写你自己的设备信息")
    _paragraph(doc, "打开 config/devices.py。每台设备是一段大括号内容。请只修改你确认过的信息，不要随便删除逗号或括号。")
    _table(doc, ["字段", "它是什么意思", "例子"], [
        ("name", "给设备起的名字，用于报告区分设备。", "34TSE-126740168"),
        ("port", "电脑识别到的串口。", "COM24"),
        ("baudrate", "串口通信速度，通常是 115200。", "115200"),
        ("user_id / license_id", "脚本执行准备命令时使用的设备目录信息。", "按实际设备填写"),
        ("sn", "设备序列号，用于给报告文件夹命名。", "126740168"),
    ], [3.3, 8.4, 5.6])
    _paragraph(doc, "每台设备的 name、port 和 sn 都必须和其他设备不同。脚本启动时会检查；如果看到“重复”提示，请回到本文件改正后再运行。只测一台设备时，就只保留那一台的配置。保存前建议先复制一份 devices.py 作为备份。", bold_lead="小建议：")

    _heading(doc, "五 开始测试")
    _number(doc, "确认设备通电、串口线连接好，并关闭其他可能占用串口的软件。")
    _number(doc, "在项目根目录运行：python main.py")
    _number(doc, "屏幕会列出即将测试的设备。确认无误后按 Enter。")
    _number(doc, "等待测试完成。不要拔线、断电或关闭窗口。")

    _heading(doc, "六 屏幕上的结果怎么看")
    _table(doc, ["看到的词", "用大白话解释", "下一步建议"], [
        ("PASS", "这一轮设备正常启动，找到了两个启动成功标志。", "正常，无需处理。"),
        ("FAIL", "串口和脚本都正常，但设备没有在限定时间内启动成功。", "看该设备的串口日志，检查设备启动问题。"),
        ("ERROR", "这次没有测成，常见原因是串口打不开、串口断开或人为中断。", "先查线、COM 口和电脑环境，不要直接认为设备坏了。"),
        ("慢启动", "设备启动了，但花的时间较长。", "关注是否偶发；连续出现时再排查。"),
    ], [2.4, 8.0, 6.9])

    _heading(doc, "七 想停止测试怎么办")
    _number(doc, "按一次 Ctrl+C。")
    _number(doc, "不要立刻关掉窗口。脚本会提示它正在通知设备线程停止。")
    _number(doc, "等看到设备线程已退出或没有新的日志继续输出后，再关闭窗口。")
    _paragraph(doc, "中断后的已完成明细会尽量保存，但本次不会生成总汇总报告，因为整场测试没有完整跑完。", bold_lead="请注意：")

    _heading(doc, "八 报告到哪里找")
    _paragraph(doc, "测试正常结束后，在项目目录下打开 设备测试结果报告 文件夹。每台设备有一个以 SN 开头的子文件夹。")
    _table(doc, ["文件", "用途"], [
        ("reboot_test_设备SN.xlsx", "每一轮的结果明细。最适合筛选 PASS、FAIL、ERROR。"),
        ("summary_设备SN.txt", "单台设备的简短统计和提前结束原因。"),
        ("serial_log_设备SN.txt", "串口原始日志。设备失败时最值得先看。"),
        ("汇总报告_时间.txt", "所有设备的文字汇总。"),
        ("HTML 报告", "适合直接打开查看整体情况和失败统计。"),
    ], [6.0, 11.3])

    _heading(doc, "九 常见问题")
    _table(doc, ["问题", "最常见原因", "怎么处理"], [
        ("串口打开失败或 Access denied", "COM 口被串口助手、IDE 或另一个脚本占用。", "关闭占用软件，拔插串口线后确认 COM 口，再重跑。"),
        ("一直等不到启动关键词", "设备没有正常启动，或关键词和设备日志不一致。", "打开 serial_log 文件，确认设备实际打印内容；必要时检查 config/settings.py。"),
        ("报告里出现 ERROR", "这是测试环境或串口链路异常，不一定是设备故障。", "先看摘要中的提前结束原因，再检查线材、供电和 COM 口。"),
        ("提示 name、port 或 sn 重复", "两条设备配置填了相同的设备名称、COM 口或序列号。", "每台设备分别填写自己的 name、port 和 sn，三项都不能重复。"),
        ("Excel 保存失败", "报告正在被 Excel 打开，或者目录没有写入权限。", "关闭正在打开的报告文件，确认磁盘空间和目录权限后重试。"),
    ], [4.3, 6.1, 6.9])

    _heading(doc, "十 脚本会怎么运行")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FLOW_PATH), width=Cm(15.6))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(p.add_run("图 1 看不懂代码也能理解的测试流程"), 9, color=(80, 80, 80))

    _paragraph(doc, "如果你只想记住一句话：先确认串口没被占用，再运行 main.py；失败先看对应设备的 serial_log 文件。", bold_lead="一句话记忆：")
    doc.save(GUIDE_DOC)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    create_flowchart()
    create_change_doc()
    create_guide_doc()
    print(CHANGE_DOC)
    print(GUIDE_DOC)
    print(FLOW_PATH)


if __name__ == "__main__":
    main()
