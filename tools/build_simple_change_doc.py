"""生成给非技术使用者阅读的简明修改说明。"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "非代码资料" / "交付文档" / "设备重启测试修改说明简单版.docx"


def set_font(run, size=11, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    tc_pr.append(node)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_text(doc, text, size=11, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    if bold_lead and text.startswith(bold_lead):
        set_font(paragraph.add_run(bold_lead), size, bold=True)
        set_font(paragraph.add_run(text[len(bold_lead):]), size)
    else:
        set_font(paragraph.add_run(text), size)
    return paragraph


def add_heading(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.paragraph_format.space_after = Pt(7)
    set_font(paragraph.add_run(text), 15, bold=True)
    return paragraph


def add_table(doc, rows, headers=("改了什么", "以前可能出现什么问题", "现在对你有什么好处")):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (4.0, 6.3, 7.0)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Cm(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        shade(cell, "245A84")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(header), 10, bold=True, color=(255, 255, 255))

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                shade(cell, "F2F6F9")
            p = cell.paragraphs[0]
            set_font(p.add_run(value), 10.5, bold=index == 0)
    return table


def create_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.35

    title_style = doc.styles["Title"]
    title_style.font.name = "Microsoft YaHei"
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.font.size = Pt(22)
    title_style.font.bold = True
    for border in title_style.element.get_or_add_pPr().findall(qn("w:pBdr")):
        title_style.element.get_or_add_pPr().remove(border)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(title.add_run("设备重启测试修改说明简单版"), 22, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(subtitle.add_run("不看代码也能明白这次改了什么"), 11, color=(90, 90, 90))

    add_text(doc, "一句话结论：脚本现在会先帮你检查配置有没有填错，测试结果不会混进旧数据，按 Ctrl+C 时也会尽量安全停下来。", bold_lead="一句话结论：")

    add_heading(doc, "一 这次改动对你最有用的地方")
    add_table(doc, [
        ("开始前检查设备信息", "两台设备如果填了同一个 COM 口、名称或 SN，可能抢同一个串口或把报告放错地方。", "脚本会直接告诉你哪里重复，不会带着错误配置开始测试。"),
        ("报告只看这一次测试", "以前生成 HTML 报告时，可能把上一次留下的旧结果也拿来显示。", "现在报告只使用本次刚生成的文件，看到的结果就是这一次的结果。"),
        ("按 Ctrl+C 能安全停止", "以前遇到串口重连或等待时，停止可能要等比较久。", "现在会尽快停止等待和重连，先保存已经完成的内容，再关闭串口。"),
        ("错误和失败分开算", "串口坏了、线松了或电脑环境出问题，容易被误认为是设备启动失败。", "报告会把真正启动失败和测试没测成分开显示，方便你判断该查设备还是查电脑。"),
    ])

    add_heading(doc, "二 你实际使用时只要记住这三件事")
    for number, text in enumerate((
        "开始前确认每台设备的 COM 口、设备名称和 SN 都不重复。",
        "测试过程中不要拔串口线；如果必须停止，按一次 Ctrl+C，然后等屏幕提示已经安全退出。",
        "测试完成后，先看本次生成的 HTML 报告；出现 ERROR 时，先检查串口线、COM 口和电脑环境。",
    ), 1):
        paragraph = doc.add_paragraph(style="List Number")
        set_font(paragraph.add_run(text), 11)

    add_heading(doc, "三 这些词是什么意思")
    add_table(doc, [
        ("PASS", "设备在规定时间内正常启动。", "正常，不需要处理。"),
        ("FAIL", "串口和脚本正常，但设备没有按要求启动。", "重点检查设备本身和启动日志。"),
        ("ERROR", "这次没有真正测成，例如串口打不开、串口断开或人为中断。", "先检查串口线、COM 口和电脑环境。"),
    ], headers=("结果", "意思", "你应该怎么做"))

    add_text(doc, "注意：这次修改没有改变设备的重启次数、启动关键词和超时时间，只是让脚本更不容易填错、结果更可靠、停止时更安全。", bold_lead="注意：")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    create_document()
    print(OUTPUT)
